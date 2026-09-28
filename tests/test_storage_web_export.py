from concurrent.futures import ThreadPoolExecutor
import json
from pathlib import Path
import sqlite3

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from irestock_demo.allocation import replenish
from irestock_demo.domain import Dataset
from irestock_demo.exporter import write_result
from irestock_demo.storage import RunRepository
from irestock_demo.web import create_app


PAYLOAD = {"rows": [{"sku": "SKU-1", "style": "STYLE-1", "store": "STORE-A",
                     "store_brand": "BRAND-A", "sales_30d": 8}], "warehouse": {"SKU-1": 8}}


def test_repository_round_trip_and_transactional_evidence(tmp_path):
    repository = RunRepository(tmp_path / "demo.sqlite3")
    result = replenish(Dataset.from_dict(PAYLOAD))
    saved = repository.save(PAYLOAD, result)
    assert repository.get(saved["run_id"]) == saved
    with repository.connect() as connection:
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
        assert connection.execute("SELECT count(*) FROM decisions").fetchone()[0] == 2


def test_evidence_failure_rolls_back_run(tmp_path):
    repository = RunRepository(tmp_path / "demo.sqlite3")
    with repository.connect() as connection:
        connection.execute("CREATE TRIGGER reject_evidence BEFORE INSERT ON decisions BEGIN SELECT RAISE(ABORT, 'forced'); END")
    with pytest.raises(sqlite3.IntegrityError):
        repository.save(PAYLOAD, replenish(Dataset.from_dict(PAYLOAD)))
    with repository.connect() as connection:
        assert connection.execute("SELECT count(*) FROM runs").fetchone()[0] == 0


def test_concurrent_results_use_independent_connections(tmp_path):
    repository = RunRepository(tmp_path / "demo.sqlite3")
    result = replenish(Dataset.from_dict(PAYLOAD))
    with ThreadPoolExecutor(max_workers=4) as executor:
        ids = list(executor.map(lambda _: repository.save(PAYLOAD, result)["run_id"], range(12)))
    assert len(set(ids)) == 12
    assert all(repository.get(run_id) for run_id in ids)


def test_api_validates_runs_and_retrieves_audit(tmp_path):
    with TestClient(create_app(tmp_path / "web.sqlite3")) as client:
        assert client.get("/health").json()["status"] == "ok"
        response = client.post("/api/run/replenishment", json=PAYLOAD)
        assert response.status_code == 200
        saved = response.json()
        assert client.get(f"/api/runs/{saved['run_id']}").json() == saved
        assert client.get("/api/runs/missing").status_code == 404
        assert client.post("/api/run/replenishment", json={"rows": PAYLOAD["rows"] * 2}).status_code == 422
        assert client.post("/api/run/unknown", json=PAYLOAD).status_code == 422


def test_excel_and_csv_preserve_null_and_escape_formula_text(tmp_path):
    payload = json.loads(json.dumps(PAYLOAD))
    payload["rows"][0].update(store="=FAKE_FORMULA", in_order=1)
    result = replenish(Dataset.from_dict(payload))
    paths = write_result(result, tmp_path, excel=True)
    assert {path.suffix for path in paths} == {".json", ".csv", ".xlsx"}
    workbook = load_workbook(tmp_path / "result.xlsx")
    sheet = workbook["audit"]
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, headers.index("store") + 1).value == "'=FAKE_FORMULA"
    assert sheet.cell(2, headers.index("target_stock") + 1).value is None
    workbook.close()


@pytest.mark.parametrize("name,mode", [("replenishment", "replenishment"), ("transfer", "transfer")])
def test_supplied_synthetic_examples_through_api(tmp_path, name, mode):
    payload = json.loads((Path(__file__).parents[1] / "examples" / f"{name}.json").read_text())
    with TestClient(create_app(tmp_path / "web.sqlite3")) as client:
        response = client.post(f"/api/run/{mode}", json=payload)
        assert response.status_code == 200
        assert len(response.json()["result"]["audit"]) == len(payload["rows"])
